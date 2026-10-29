"""Prism Intake: anything a company hands over is read by code, understood, checked, and only drafted."""

from __future__ import annotations

import asyncio
import io
import json
import re
import time
import zipfile

import pytest

from verytas.intake import fetch, readers, tables
from verytas.intake.tables import Dedupe, DropHeaderRows, FactPlan, Rename

M = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def xlsx(rows: list[list[str | float]], sheet="Pricing") -> bytes:
    strings: list[str] = []

    def cell(ref: str, v) -> str:
        if isinstance(v, (int, float)):
            return f'<c r="{ref}"><v>{v}</v></c>'
        strings.append(v)
        return f'<c r="{ref}" t="s"><v>{len(strings) - 1}</v></c>'

    body = "".join(f'<row r="{i + 1}">' + "".join(cell(f"{chr(65 + j)}{i + 1}", v) for j, v in enumerate(r)) + "</row>"
                   for i, r in enumerate(rows))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("xl/workbook.xml", f'<workbook xmlns="{M}" xmlns:r="{R}"><sheets><sheet name="{sheet}" '
                                      f'sheetId="1" r:id="rId1"/></sheets></workbook>')
        z.writestr("xl/_rels/workbook.xml.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/'
                   '2006/relationships"><Relationship Id="rId1" Target="worksheets/sheet1.xml"/></Relationships>')
        z.writestr("xl/worksheets/sheet1.xml", f'<worksheet xmlns="{M}"><sheetData>{body}</sheetData></worksheet>')
        z.writestr("xl/sharedStrings.xml", f'<sst xmlns="{M}">' + "".join(f"<si><t>{s}</t></si>" for s in strings)
                   + "</sst>")
    return buf.getvalue()


def pptx(slides: list[str]) -> bytes:
    a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("ppt/presentation.xml", "<p/>")
        for i, text in enumerate(slides, 1):
            z.writestr(f"ppt/slides/slide{i}.xml", f'<sld xmlns:a="{a}"><a:p><a:r><a:t>{text}</a:t></a:r></a:p></sld>')
    return buf.getvalue()


def test_slide_tables_and_speaker_notes_and_word_tables_are_read_as_rows():
    a, pn = "http://schemas.openxmlformats.org/drawingml/2006/main", "http://schemas.openxmlformats.org/presentationml/2006/main"
    rel = "http://schemas.openxmlformats.org/package/2006/relationships"
    cell = lambda t: f"<a:tc><a:txBody><a:p><a:r><a:t>{t}</a:t></a:r></a:p></a:txBody></a:tc>"  # noqa: E731
    row = lambda *c: "<a:tr>" + "".join(cell(x) for x in c) + "</a:tr>"  # noqa: E731
    ph = lambda kind, t: (f'<p:sp><p:nvSpPr><p:nvPr><p:ph type="{kind}"/></p:nvPr></p:nvSpPr>'  # noqa: E731
                          f"<p:txBody><a:p><a:r><a:t>{t}</a:t></a:r></a:p></p:txBody></p:sp>")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("ppt/slides/slide1.xml", f'<p:sld xmlns:a="{a}" xmlns:p="{pn}"><a:p><a:r><a:t>Channel limits</a:t>'
                   f'</a:r></a:p><a:tbl>{row("Channel", "Limit")}{row("instagram", "2200")}</a:tbl></p:sld>')
        z.writestr("ppt/slides/_rels/slide1.xml.rels", f'<Relationships xmlns="{rel}"><Relationship Id="r1" Type='
                   '"http://x/notesSlide" Target="../notesSlides/notesSlide1.xml"/></Relationships>')
        z.writestr("ppt/notesSlides/notesSlide1.xml", f'<p:notes xmlns:a="{a}" xmlns:p="{pn}">{ph("sldNum", "1")}'
                   f'{ph("body", "Say we are rated 4.8 by parents")}</p:notes>')
    deck = readers.read(buf.getvalue(), readers.PPTX)
    assert deck.pages[0].startswith("Channel limits")
    assert deck.pages[0] == ("Channel limits\n[Slide 1 table 1]\nChannel | Limit\ninstagram | 2200\n"
                             "[Speaker notes] Say we are rated 4.8 by parents")
    assert deck.tables[0].rows == [["instagram", "2200"]] and deck.tables[0].cell(0, 1) == "Slide 1 table 1!B2"
    # A merge across two columns keeps its placeholder cell, so the column after it stays in place.
    merged = (f'<p:sld xmlns:a="{a}" xmlns:p="{pn}"><a:tbl>{row("Plan", "Price", "Months")}<a:tr>'
              '<a:tc gridSpan="2"><a:txBody><a:p><a:r><a:t>Group</a:t></a:r></a:p></a:txBody></a:tc>'
              f'<a:tc hMerge="1"><a:txBody><a:p/></a:txBody></a:tc>{cell("6")}</a:tr></a:tbl></p:sld>')
    one = io.BytesIO()
    with zipfile.ZipFile(one, "w") as z:
        z.writestr("ppt/slides/slide1.xml", merged)
    assert readers.read(one.getvalue(), readers.PPTX).tables[0].rows == [["Group", "", "6"]]

    from docx import Document
    doc = Document()
    doc.add_paragraph("Words we use")
    t = doc.add_table(rows=3, cols=2)
    for r, (say, avoid) in enumerate([("Say", "Instead of"), ("learners", "students"), ("course", "package")]):
        t.cell(r, 0).text, t.cell(r, 1).text = say, avoid
    out = io.BytesIO()
    doc.save(out)
    word = readers.read(out.getvalue(), readers.DOCX)
    assert "learners\nstudents" in word.text  # the text the review reads is unchanged
    assert word.tables[0].header == ["Say", "Instead of"] and word.tables[0].rows == [["learners", "students"],
                                                                                      ["course", "package"]]


def test_word_table_cells_read_what_the_reader_sees_within_limits():
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

    def docx(body: str) -> bytes:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as z:
            z.writestr("word/document.xml", f'<w:document xmlns:w="{w}"><w:body>{body}</w:body></w:document>')
        return buf.getvalue()

    def tc(inner: str, span: int = 1) -> str:
        props = f'<w:tcPr><w:gridSpan w:val="{span}"/></w:tcPr>' if span > 1 else ""
        return f"<w:tc>{props}<w:p>{inner}</w:p></w:tc>"

    run = lambda t: f"<w:r><w:t>{t}</w:t></w:r>"  # noqa: E731
    changed = ("<w:del><w:r><w:delText>299</w:delText></w:r></w:del><w:ins>" + run("349") + "</w:ins>")
    body = (f"<w:tbl><w:tr>{tc(run('Plan'))}{tc(run('Price'))}{tc(run('Months'))}</w:tr>"
            f"<w:tr>{tc(run('Pro'))}{tc(changed)}{tc(run('12'))}</w:tr>"
            f"<w:tr>{tc(run('Group offer'), span=2)}{tc(run('6'))}</w:tr></w:tbl>")
    [t] = readers.read(docx(body), readers.DOCX).tables
    assert t.rows == [["Pro", "349", "12"], ["Group offer", "", "6"]]  # 299 was deleted; the merge keeps Months
    tab = run("learners") + "<w:r><w:tab/></w:r>" + run("always")
    tabbed = f"<w:tbl><w:tr>{tc(run('Say'))}{tc(run('Why'))}</w:tr><w:tr>{tc(tab)}{tc(run('warmer'))}</w:tr></w:tbl>"
    assert readers.read(docx(tabbed), readers.DOCX).tables[0].rows == [["learners always", "warmer"]]
    wide = "<w:tbl>" + ("<w:tr>" + "".join(tc(run(f"h{i}")) for i in range(500)) + "</w:tr>") * 3 + "</w:tbl>"
    [t] = readers.read(docx(wide), readers.DOCX).tables
    assert len(t.header) == readers.MAX_COLS and all(len(r) == readers.MAX_COLS for r in t.rows)


def test_a_bad_notes_page_costs_its_own_words_only():
    a, pn = "http://schemas.openxmlformats.org/drawingml/2006/main", "http://schemas.openxmlformats.org/presentationml/2006/main"
    rel = "http://schemas.openxmlformats.org/package/2006/relationships"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for i in (1, 2):
            z.writestr(f"ppt/slides/slide{i}.xml", f'<p:sld xmlns:a="{a}" xmlns:p="{pn}"><a:p><a:r><a:t>Slide {i}</a:t>'
                       "</a:r></a:p></p:sld>")
            z.writestr(f"ppt/slides/_rels/slide{i}.xml.rels", f'<Relationships xmlns="{rel}"><Relationship Id="r1" '
                       f'Type="http://x/notesSlide" Target="/ppt/notesSlides/notesSlide{i}.xml"/></Relationships>')
        z.writestr("ppt/notesSlides/notesSlide1.xml", "x" * (readers.MAX_XML + 1))
        z.writestr("ppt/notesSlides/notesSlide2.xml", f'<p:notes xmlns:a="{a}" xmlns:p="{pn}"><p:sp><p:nvSpPr><p:nvPr>'
                   '<p:ph type="body"/></p:nvPr></p:nvSpPr><p:txBody><a:p><a:r><a:t>Mention the 4.8 rating</a:t>'
                   "</a:r></a:p></p:txBody></p:sp></p:notes>")
    deck = readers.read(buf.getvalue(), readers.PPTX)
    assert deck.pages == ["Slide 1", "Slide 2\n[Speaker notes] Mention the 4.8 rating"]
    assert any("speaker notes of slide 1" in n for n in deck.notes)


def test_office_files_and_text_formats_are_read_with_addresses():
    raw = xlsx([["Kitepath price list", ""], ["Plan", "Price (Rs)", "From"], ["Basic", 299, 46023],
                ["Pro", 599, 46023], ["", "", ""]])
    assert readers.office_kind(raw) == readers.XLSX
    r = readers.read(raw, readers.XLSX)
    [t] = r.tables
    # The title row above the table is not taken for the header.
    assert t.header == ["Plan", "Price (Rs)", "From"] and t.rows[0][:2] == ["Basic", "299"]
    assert t.cell(1, 1) == "Pricing!B4"
    assert readers.read(pptx(["Always say Kitepath", "Never mock other apps"]), readers.PPTX).pages[1] \
        == "Never mock other apps"
    assert readers.sniff_text(b"plan;price\nBasic;299\nPro;599\n") == "text/plain"
    assert readers.sniff_text(b"plan,price\nBasic,299\nPro,599\n") == "text/csv"
    data = readers.read(json.dumps({"data": [{"plan": "Basic", "price": 299}]}).encode(), "application/json", "api")
    assert data.tables[0].header == ["plan", "price"]
    assert readers.sniff_text(b"\x00\x01\x02binary") is None
    # Entities are never expanded: a billion-laughs part reads as nothing, not as gigabytes.
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w") as z:
        z.writestr("xl/workbook.xml", '<!DOCTYPE x [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;&a;">]>'
                                      f'<workbook xmlns="{M}"><sheets/></workbook>')
        z.writestr("xl/_rels/workbook.xml.rels", "<Relationships/>")
    assert readers.read(bomb.getvalue(), readers.XLSX).tables == []


def test_the_model_plans_cleaning_from_a_menu_and_code_makes_one_fact_per_row():
    t = readers.Table("Pricing", ["Plan", "Price", "Note"],
                      [["Basic", "₹299/month", ""], ["Pro", "₹599/month", ""], ["Pro", "₹599/month", ""],
                       ["All prices include GST", "", ""], ["Annual", "4,999", "per year"]], header_row=2)
    cleaned, origin, skipped = tables.apply(t, [DropHeaderRows(column="Plan", pattern="include GST"),
                                                Dedupe(columns=["Plan"]), Rename(column="Nope", to="x")])
    assert [r[0] for r in cleaned.rows] == ["Basic", "Pro", "Annual"] and origin == [0, 1, 4]
    assert skipped == ["rename: no column named 'Nope'"]
    rows, problems = tables.facts(cleaned, FactPlan(statement="The {Plan} plan costs {Price}", subject="Plan",
                                                    value="Price"), origin, "prices.xlsx")
    assert rows[0]["payload"]["value"] == 299 and rows[0]["payload"]["unit"] == "INR"
    assert rows[0]["provenance"]["cell"] == "Pricing!B3" and rows[2]["provenance"]["cell"] == "Pricing!B7"
    assert tables.facts(cleaned, FactPlan(statement="{Missing}"), origin, "x")[1][0].startswith("the plan names")
    assert tables.number("12%") == (12, "%") and tables.number("no") is None
    assert tables.as_date("31/12/2026").isoformat() == "2026-12-31"


def test_a_live_source_must_be_public_https():
    for url, why in (("http://example.com/a.csv", "only https"), ("https://user:pw@example.com/x", "no password")):
        with pytest.raises(fetch.FetchRefused, match=why):
            fetch.check_url(url)
    with pytest.raises(fetch.FetchRefused, match="private network"):
        asyncio.run(fetch.fetch("https://127.0.0.1/prices.csv"))


def _wait(client, h, batch_id, timeout=20.0) -> dict:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        b = client.get(f"/intake/{batch_id}", headers=h).json()
        if b["status"] in ("done", "failed") and all(i["status"] not in ("queued", "reading") for i in b["items"]):
            return b
        time.sleep(0.1)
    raise AssertionError(f"intake {batch_id} did not finish: {b}")


def test_an_admin_hands_over_material_and_approves_only_what_reads_right(client):
    h = {"Authorization": "Bearer " + client.post("/bootstrap", json={
        "workspace_name": "Kitepath", "admin_email": "a@k.io", "admin_name": "Meera"}).json()["token"]}
    up = client.post("/intake/files", headers=h, files={"file": ("prices.xlsx", xlsx(
        [["Plan", "Price (Rs)"], ["Basic", 299], ["Pro", 599], ["Pro", 649]]))}).json()
    assert up["media_type"] == readers.XLSX
    from PIL import Image

    png = io.BytesIO()
    Image.new("RGB", (400, 300), (30, 90, 200)).save(png, "PNG")
    pic = client.post("/intake/files", headers=h, files={"file": ("chart.png", png.getvalue())}).json()
    started = client.post("/intake", headers=h, json={"note": None, "items": [
        {"blob_id": up["id"]},
        {"text": "Always write Kitepath with a capital K. Never mock other learning apps. Project Falcon is our "
                 "unreleased tutor app.", "name": "Brand book extract"},
        {"blob_id": pic["id"]}]})
    assert started.status_code == 201, started.text
    b = _wait(client, h, started.json()["id"])
    by_name = {i["name"]: i for i in b["items"]}
    assert by_name["prices.xlsx"]["status"] == "understood"
    # A picture with nothing said about it is a question, and adds nothing until answered.
    assert by_name["chart.png"]["status"] == "needs_context" and by_name["chart.png"]["questions"]
    assert not [d for d in b["drafts"] if d["item_id"] == by_name["chart.png"]["id"]]
    facts = [d for d in b["drafts"] if d["target"] == "fact"]
    assert {d["payload"]["value"] for d in facts} == {299, 599, 649}
    assert all(d["provenance"]["cell"].startswith("Pricing!B") for d in facts)
    # Two prices for Pro in one sheet: flagged, so bulk approval leaves it for a person.
    flagged = [d for d in facts if any(i["kind"] == "conflict" for i in d["issues"])]
    assert [d["payload"]["value"] for d in flagged] == [649]
    assert {d["target"] for d in b["drafts"]} >= {"rule", "term", "fact"}
    assert any(d["payload"]["term"] == "Falcon" or "Falcon" in d["payload"]["term"] for d in b["drafts"]
               if d["target"] == "term")
    # Nothing is in the library yet.
    assert client.get("/facts", headers=h).json() == []
    res = client.post(f"/intake/{b['id']}/decide", headers=h, json={"ids": [d["id"] for d in facts]}).json()
    assert set(res["skipped"]) == {flagged[0]["id"]} and len(res["done"]) == 2
    assert sorted(f["value"] for f in client.get("/facts", headers=h).json()) == [299, 599]
    one = client.post(f"/intake/drafts/{flagged[0]['id']}", headers=h, json={"action": "reject"}).json()
    assert one["status"] == "rejected"
    # Answering the question reads the picture again with the answer.
    ans = client.post(f"/intake/items/{by_name['chart.png']['id']}/answer", headers=h,
                      json={"context": "Our brand blue, as a colour swatch for designers."})
    assert ans.status_code == 200
    again = _wait(client, h, b["id"])
    assert {i["name"]: i for i in again["items"]}["chart.png"]["status"] == "understood"
    # Brand admins only; and a creator can't even see that an intake exists.
    creator = client.post("/users", headers=h, json={"email": "c@k.io", "name": "Cy", "role": "creator"}).json()
    assert client.get(f"/intake/{b['id']}", headers={"Authorization": f"Bearer {creator['token']}"}).status_code == 403


def test_hostile_or_odd_files_stay_bounded_and_addresses_stay_true():
    # One long shared string in every cell: each cell is cut, so text and preview stay small.
    raw = xlsx([["Plan", "Note"]] + [[f"P{i}", "BIG"] for i in range(300)])
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    # Every "Note" cell points at shared string 1, which is now 2 MB.
    big = "x" * 2_000_000
    parts["xl/sharedStrings.xml"] = f'<sst xmlns="{M}"><si><t>Plan</t></si><si><t>{big}</t></si></sst>'.encode()
    sheet = parts["xl/worksheets/sheet1.xml"].decode()
    parts["xl/worksheets/sheet1.xml"] = re.sub(r'(<c r="B\d+" t="s"><v>)\d+', r"\g<1>1", sheet).encode()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in parts.items():
            z.writestr(n, b)
    r = readers.read(buf.getvalue(), readers.XLSX)
    assert len(r.text) <= readers.MAX_TEXT and len(json.dumps(r.preview())) < 200_000
    assert all(len(c) <= readers.MAX_CELL for row in r.tables[0].rows for c in row)
    # A blank row inside the table keeps the rows below it at their real addresses.
    t = readers.read(b"Plan,Price\nBasic,299\n,\nPro,599\n", "text/csv", "p.csv").tables[0]
    assert t.cell(1, 1) == "p!B4"
    cleaned, origin, _ = tables.apply(t, [Dedupe(columns=["Plan"])])
    rows, _ = tables.facts(cleaned, FactPlan(statement="{Plan}: {Price}", subject="Plan", value="Price"), origin, "p")
    assert [f["provenance"]["cell"] for f in rows] == ["p!B2", "p!B4"]
    # Only spreadsheets have serial dates, and "Price from" is a price.
    t = readers.read(b"Plan,Price from\nBasic,25000\nPro,45000\n", "text/csv", "p.csv").tables[0]
    assert [r[1] for r in t.rows] == ["25000", "45000"]
    # Prose that starts with a bracket is text, not broken JSON.
    note = b"[Draft] Always write Kitepath with a capital K.\nNever mock rivals."
    assert readers.sniff_text(note) == "text/plain"
    assert "capital K" in readers.read(note, "application/json", "a").text
    # A picture over the pixel budget is skipped before it is decoded.
    from PIL import Image

    from verytas.media.files import MAX_PIXELS

    side = int(MAX_PIXELS ** 0.5) + 100
    buf = io.BytesIO()
    Image.new("1", (side, side)).save(buf, "PNG")
    assert readers._as_png(buf.getvalue()) is None


def test_sheet_dates_come_from_cell_styles():
    raw = xlsx([["Plan", "Price from", "Launch"], ["Basic", 25000, 46023]])
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    parts["xl/worksheets/sheet1.xml"] = parts["xl/worksheets/sheet1.xml"].replace(
        b'<c r="C2"><v>46023</v></c>', b'<c r="C2" s="1"><v>46023</v></c>')
    parts["xl/styles.xml"] = (f'<styleSheet xmlns="{M}"><cellXfs><xf numFmtId="0"/><xf numFmtId="14"/></cellXfs>'
                              '</styleSheet>').encode()
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for n, b in parts.items():
            z.writestr(n, b)
    [t] = readers.read(out.getvalue(), readers.XLSX).tables
    assert t.rows[0] == ["Basic", "25000", "2026-01-01"]


def test_approvals_are_claimed_once_and_a_refused_one_does_not_break_the_rest(client):
    h = {"Authorization": "Bearer " + client.post("/bootstrap", json={
        "workspace_name": "Kitepath", "admin_email": "a@k.io", "admin_name": "Meera"}).json()["token"]}
    text = ("Always write Kitepath with a capital K. Never mock other learning apps. Project Falcon is our "
            "unreleased tutor app.")

    def batch(t: str) -> dict:
        started = client.post("/intake", headers=h, json={"items": [{"text": t, "name": "Brand book"}]})
        assert started.status_code == 201, started.text
        return _wait(client, h, started.json()["id"])

    first = batch(text)
    rules = [d for d in first["drafts"] if d["target"] == "rule"]
    # A brand book's rules reach public posts, as the brand-book importer's do.
    assert rules and all(d["payload"]["tier"] == "public" for d in rules)
    # The same material again proposes nothing the first batch already drafted.
    again = batch(text)
    assert not [d for d in again["drafts"] if d["target"] in ("rule", "term")]
    assert any(i["kind"] == "known" for i in again["items"][0]["issues"])

    # One draft, two approvals: the second finds it taken and creates nothing.
    one = rules[0]
    assert client.post(f"/intake/drafts/{one['id']}", headers=h, json={"action": "approve"}).status_code == 200
    assert client.post(f"/intake/drafts/{one['id']}", headers=h, json={"action": "approve"}).status_code == 409
    assert sum(r["text"] == one["payload"]["text"] for r in client.get("/rules", headers=h).json()) == 1

    # The library refuses one draft of a bulk approval (the codename was added by hand meanwhile): that one
    # is marked and left pending, the rest go in.
    term = next(d for d in first["drafts"] if d["target"] == "term")
    assert client.post("/confidential-terms", headers=h, json={"term": term["payload"]["term"]}).status_code == 201
    rest = [d["id"] for d in first["drafts"] if d["status"] == "pending" and d["id"] != one["id"]
            and not d["issues"]]
    res = client.post(f"/intake/{first['id']}/decide", headers=h, json={"ids": rest})
    assert res.status_code == 200, res.text
    assert term["id"] in res.json()["skipped"] and res.json()["done"]
    after = {d["id"]: d for d in client.get(f"/intake/{first['id']}", headers=h).json()["drafts"]}
    assert after[term["id"]]["status"] == "pending"
    assert any(i["kind"] == "rejected_by_library" for i in after[term["id"]]["issues"])


def test_a_signed_link_never_shows_its_key():
    from verytas.intake import pipeline

    assert pipeline.public_source("https://u:p@sheets.example:8443/d/x.csv?X-Amz-Signature=abc#g") \
        == "https://sheets.example:8443/d/x.csv"


def test_a_database_refusal_puts_the_draft_back_and_a_typed_name_is_kept(client, monkeypatch):
    from sqlalchemy.exc import OperationalError

    from verytas.api import intake as api

    h = {"Authorization": "Bearer " + client.post("/bootstrap", json={
        "workspace_name": "Kitepath", "admin_email": "a@k.io", "admin_name": "Meera"}).json()["token"]}
    started = client.post("/intake", headers=h, json={"items": [
        {"text": "Always write Kitepath with a capital K. Project Falcon is our unreleased tutor app.",
         "name": "Brand book"}]})
    b = _wait(client, h, started.json()["id"])
    draft = b["drafts"][0]

    async def locked(*_a, **_k):
        raise OperationalError("insert", {}, Exception("database is locked"))

    monkeypatch.setattr(api, "_apply", locked)
    assert client.post(f"/intake/drafts/{draft['id']}", headers=h, json={"action": "approve"}).status_code == 409
    monkeypatch.undo()
    # Not stuck half way: it can be approved now.
    assert client.post(f"/intake/drafts/{draft['id']}", headers=h, json={"action": "approve"}).status_code == 200
    assert api._link_name("https://x.example/a/prices.csv?sig=1") == "prices.csv"
