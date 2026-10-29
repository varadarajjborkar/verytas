"""Builds a brand-kit corpus with known content and locations, plus hostile files. Writes corpus/truth.json."""

import io
import json
import random
import sys
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Image as RLImage
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table

OUT = Path(sys.argv[1]).resolve()  # the XXE entity needs an absolute file path
OUT.mkdir(exist_ok=True)
random.seed(7)
SUBJ = [
    "the logo",
    "a headline",
    "a caption",
    "our mascot Kit",
    "a price",
    "a student quote",
    "the tagline",
    "a course name",
    "a discount",
    "a testimonial",
    "a hashtag",
    "the brand colour",
    "an emoji",
    "a CTA",
]
VERB = ["must always", "should never", "may only", "must not", "should"]
OBJ = [
    "appear on a busy photo background",
    "promise exam results",
    "mention a rival platform by name",
    "use more than two exclamation marks",
    "be shown below 24 pixels",
    "name a student without consent",
    "quote a placement rate without a source",
    "be translated without review",
    "run past the offer end date",
    "sit within 12 mm of the edge",
    "be written in title case",
    "use the word guaranteed",
    "carry the registered mark",
    "link to the official store",
]


def sentence(tag):
    return f"{random.choice(SUBJ).capitalize()} {random.choice(VERB)} {random.choice(OBJ)} (rule {tag})."


truth = {"files": {}}
st = getSampleStyleSheet()

# 1. PDF: 6 pages of rules, a table on page 4, a running footer on every page.
pdf_sents, story = [], []
table_rows = [["Term", "Use", "Avoid"]] + [[f"Term{i}x", f"use-{i}-ok", f"avoid-{i}-no"] for i in range(1, 7)]
for page in range(1, 7):
    story.append(Paragraph(f"Section {page}: brand rules", st["Heading2"]))
    for k in range(1, 5):
        s = sentence(f"P{page}.{k}")
        pdf_sents.append({"text": s, "page": page})
        story.append(Paragraph(s, st["BodyText"]))
    if page == 4:
        story.append(Spacer(1, 8))
        story.append(Table(table_rows))
    if page < 6:
        story.append(PageBreak())


def footer(c, d):
    c.setFont("Helvetica", 8)
    c.drawString(40, 20, "Kitepath brand guide, internal")


SimpleDocTemplate(str(OUT / "guidelines.pdf"), pagesize=A4).build(story, onFirstPage=footer, onLaterPages=footer)
truth["files"]["guidelines.pdf"] = {"sentences": pdf_sents, "tables": [{"rows": table_rows, "page": 4}]}

# 2. Word: headings, paragraphs, a bullet list, a table, a header.
doc = Document()
doc.sections[0].header.paragraphs[0].text = "Kitepath voice guide"
doc_sents = []
for h in range(1, 5):
    doc.add_heading(f"Voice part {h}", level=1)
    for k in range(1, 5):
        s = sentence(f"W{h}.{k}")
        doc_sents.append({"text": s})
        doc.add_paragraph(s)
    s = sentence(f"W{h}.b")
    doc_sents.append({"text": s})
    doc.add_paragraph(s, style="List Bullet")
wrows = [["Say", "Instead of", "Why"]] + [[f"say{i}word", f"not{i}word", f"reason{i}text"] for i in range(1, 6)]
t = doc.add_table(rows=len(wrows), cols=3)
for r, row in enumerate(wrows):
    for c, v in enumerate(row):
        t.cell(r, c).text = v
doc.save(OUT / "voice.docx")
truth["files"]["voice.docx"] = {"sentences": doc_sents, "tables": [{"rows": wrows}]}

# 3. Excel: two sheets, a gap row, dates and numbers; every cell's address is known.
wb = Workbook()
ws = wb.active
ws.title = "Approved claims"
claims = [["Claim", "Source", "Owner", "Score"]] + [
    [f"Claim number {i} about courses", f"https://kitepath.example/src{i}", f"owner{i}", i * 3] for i in range(1, 21)
]
cells = []
for r, row in enumerate(claims):
    line = r + 1 if r < 11 else r + 2  # a blank row after row 11
    for c, v in enumerate(row):
        ws.cell(row=line, column=c + 1, value=v)
        cells.append({"sheet": "Approved claims", "addr": f"{'ABCD'[c]}{line}", "value": str(v)})
ws2 = wb.create_sheet("Pricing")
prices = [["Plan", "INR", "Months"]] + [[f"Plan{i}", 499 * i, i * 3] for i in range(1, 9)]
for r, row in enumerate(prices):
    for c, v in enumerate(row):
        ws2.cell(row=r + 1, column=c + 1, value=v)
        cells.append({"sheet": "Pricing", "addr": f"{'ABC'[c]}{r + 1}", "value": str(v)})
wb.save(OUT / "claims.xlsx")
truth["files"]["claims.xlsx"] = {
    "cells": cells,
    "tables": [{"rows": [[str(v) for v in r] for r in claims]}, {"rows": [[str(v) for v in r] for r in prices]}],
}

# 4. PowerPoint: 6 slides, bullets, speaker notes, a table on slide 3.
prs = Presentation()
ppt_sents = []
for n in range(1, 7):
    sl = prs.slides.add_slide(prs.slide_layouts[1])
    sl.shapes.title.text = f"Campaign slide {n}"
    body = sl.placeholders[1].text_frame
    for k in range(1, 4):
        s = sentence(f"S{n}.{k}")
        ppt_sents.append({"text": s, "page": n})
        (body.paragraphs[0] if k == 1 else body.add_paragraph()).text = s
    note = sentence(f"N{n}")
    ppt_sents.append({"text": note, "page": n, "notes": True})
    sl.notes_slide.notes_text_frame.text = note
    if n == 3:
        prow = [["Channel", "Limit"], ["instagram", "2200chars"], ["linkedin", "3000chars"]]
        tb = sl.shapes.add_table(3, 2, Inches(1), Inches(5), Inches(4), Inches(1)).table
        for r, row in enumerate(prow):
            for c, v in enumerate(row):
                tb.cell(r, c).text = v
prs.save(OUT / "deck.pptx")
truth["files"]["deck.pptx"] = {"sentences": ppt_sents, "tables": [{"rows": prow, "page": 3}]}

# 5. CSV.
crow = [["code", "term", "status"]] + [
    [f"T{i:03d}", f"codename{i}", random.choice(["banned", "allowed"])] for i in range(1, 41)
]
(OUT / "terms.csv").write_text("\n".join(",".join(r) for r in crow) + "\n")
truth["files"]["terms.csv"] = {"tables": [{"rows": crow}]}

# 6. A scanned page: the text is only in a picture.
img = Image.new("RGB", (1200, 400), "white")
ImageDraw.Draw(img).text(
    (40, 150), "Scanned rule: never crop the logo", fill="black", font=ImageFont.load_default(size=48)
)
buf = io.BytesIO()
img.save(buf, "PNG")
buf.seek(0)
SimpleDocTemplate(str(OUT / "scanned.pdf"), pagesize=A4).build([RLImage(buf, width=480, height=160)])
truth["files"]["scanned.pdf"] = {"scanned": True}


# Hostile files, built from a clean Word file by rewriting its document part.
def rewrite_docx(name, doctype, text):
    src = zipfile.ZipFile(OUT / "voice.docx")
    out = zipfile.ZipFile(OUT / name, "w", zipfile.ZIP_DEFLATED)
    for item in src.infolist():
        data = src.read(item)
        if item.filename == "word/document.xml":
            xml = data.decode()
            head, rest = xml.split("?>", 1)
            body = rest.replace("Voice part 1", text, 1)
            data = (head + "?>" + doctype + body).encode()
        out.writestr(item, data)
    out.close()


canary = OUT.parent / "canary.txt"
canary.write_text("CANARY-7731-SECRET\n")
rewrite_docx("xxe.docx", f'<!DOCTYPE w:document [<!ENTITY xxe SYSTEM "file://{canary}">]>', "&xxe;")
lol = "".join(f'<!ENTITY l{i} "{f"&l{i - 1};" * 10 if i else "lol"}">' for i in range(9))
rewrite_docx("laughs.docx", f"<!DOCTYPE w:document [{lol}]>", "&l8;")
# A spreadsheet whose second sheet unpacks to 300 MB from about 300 KB (openpyxl keeps text in the sheets).
src = zipfile.ZipFile(OUT / "claims.xlsx")
out = zipfile.ZipFile(OUT / "zipbomb.xlsx", "w", zipfile.ZIP_DEFLATED)
for item in src.infolist():
    data = src.read(item)
    if item.filename == "xl/worksheets/sheet2.xml":
        head, tail = data.split(b"</row>", 1)
        with out.open(item.filename, "w", force_zip64=True) as f:
            f.write(head + b'<c r="D1" t="inlineStr"><is><t>')
            for _ in range(300):
                f.write(b"a" * 1_000_000)
            f.write(b"</t></is></c></row>" + tail)
    else:
        out.writestr(item, data)
out.close()
truth["hostile"] = {"xxe.docx": "CANARY-7731-SECRET", "laughs.docx": None, "zipbomb.xlsx": None}
(OUT / "truth.json").write_text(json.dumps(truth, indent=1))
print("built", sorted(p.name for p in OUT.iterdir()))
